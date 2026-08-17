from pathlib import Path

path = Path("main.py")
text = path.read_text(encoding="utf-8")


def once(old, new, label):
    global text
    if old not in text:
        raise RuntimeError(f"Не найден участок: {label}")
    text = text.replace(old, new, 1)


once('__version__ = "1.4.0"', '__version__ = "1.5.0"', "версия")

once(
    'def create_document(callback):',
    'def create_document(callback, filename="Отчет_сравнения_актов.xlsx"):',
    "имя сохраняемого Excel",
)
once(
    'intent.putExtra(Intent.EXTRA_TITLE, "Отчет_сравнения_актов.xlsx")',
    'intent.putExtra(Intent.EXTRA_TITLE, filename)',
    "имя файла Android",
)

FNS_FUNCTIONS = r'''FNS_RECONCILIATION_ORDER = "Приказ ФНС России от 13.05.2022 № ЕД-7-26/405@"
FNS_FORMAT_VERSION = "5.01"
FNS_FORMAT_PART = "972"
FNS_RECIPIENT_KND = "1110333"


def _fns_date(value):
    value = smart_date(value)
    if not value:
        return ""
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m.%Y")
    except Exception:
        return str(value)


def _fns_key_identity(key):
    key = str(key or "")
    parts = [part.strip() for part in key.split("|") if part.strip()]
    doc_type = parts[0] if len(parts) >= 1 and not parts[0].startswith("строка ") else ""
    doc_number = parts[1] if len(parts) >= 2 else ""
    doc_date = parts[2] if len(parts) >= 3 else ""
    if not doc_type:
        doc_type = smart_type(key)
    if not doc_number and not key.startswith("строка "):
        doc_number = smart_number(key)
    if not doc_date:
        doc_date = smart_date(key)
    return doc_type or "Документ", doc_number or "б/н", _fns_date(doc_date)


def _fns_row_amounts(row, cols):
    debit = number(row.get(cols.get("debit"))) if cols.get("debit") else None
    credit = number(row.get(cols.get("credit"))) if cols.get("credit") else None
    return debit, credit


def _fns_missing_rows(items, cols, source_name, direction):
    rows = []
    for item in items:
        doc_type = item.get("Авто: тип") or smart_type(" | ".join(str(v) for v in item.values())) or "Документ"
        doc_number = item.get("Авто: номер") or smart_number(" | ".join(str(v) for v in item.values())) or "б/н"
        doc_date = _fns_date(item.get("Авто: дата") or smart_date(" | ".join(str(v) for v in item.values())))
        debit, credit = _fns_row_amounts(item, cols)
        amount = debit if debit is not None else credit if credit is not None else 0.0
        amount = float(amount or 0.0) * direction
        rows.append({
            "НаимДок": doc_type,
            "НомДок": doc_number,
            "ДатаДок": doc_date,
            "ИтогоРасх": amount,
            "ДатаОпер": doc_date,
            "НаимОпер": f"Документ присутствует только в {source_name}",
            "СумДебет": debit,
            "СумКредит": credit,
            "ДопИнф": item.get("Рекомендация", "Проверить наличие документа у второй стороны."),
        })
    return rows


def _fns_discrepancy_rows(differences, only1, only2, cols1, cols2):
    result = []
    grouped = {}
    for diff in differences:
        key = str(diff[0])
        doc_type, doc_number, doc_date = _fns_key_identity(key)
        group = grouped.setdefault(key, {
            "НаимДок": doc_type,
            "НомДок": doc_number,
            "ДатаДок": doc_date,
            "ИтогоРасх": 0.0,
            "items": [],
        })
        field = str(diff[1]).upper()
        value1 = diff[4] if len(diff) > 4 else ""
        value2 = diff[5] if len(diff) > 5 else ""
        delta = number(diff[6]) if len(diff) > 6 else None
        recommendation = diff[8] if len(diff) > 8 else "Проверить исходные документы."

        if field == "DOCUMENT":
            group["НаимДок"] = smart_type(value2) or smart_type(value1) or group["НаимДок"]
            group["НомДок"] = smart_number(value2) or smart_number(value1) or group["НомДок"]
        elif field == "DATE":
            group["ДатаДок"] = _fns_date(value2) or _fns_date(value1) or group["ДатаДок"]
        elif field in ("DEBIT", "CREDIT") and delta is not None:
            group["ИтогоРасх"] += float(delta)

        recipient_amount = number(value2)
        sum_debit = recipient_amount if field == "DEBIT" else None
        sum_credit = recipient_amount if field == "CREDIT" else None
        group["items"].append({
            "ДатаОпер": group["ДатаДок"],
            "НаимОпер": f"Расхождение {field}: Акт 1 = {value1}; Акт 2 = {value2}",
            "СумДебет": sum_debit,
            "СумКредит": sum_credit,
            "ДопИнф": str(recommendation)[:500],
        })

    for group in grouped.values():
        if not group["items"]:
            group["items"].append({
                "ДатаОпер": group["ДатаДок"], "НаимОпер": "Расхождение",
                "СумДебет": None, "СумКредит": None, "ДопИнф": "Проверить исходные документы.",
            })
        for item in group["items"]:
            result.append({
                "НаимДок": group["НаимДок"] or "Документ",
                "НомДок": group["НомДок"] or "б/н",
                "ДатаДок": group["ДатаДок"],
                "ИтогоРасх": group["ИтогоРасх"],
                **item,
            })

    result.extend(_fns_missing_rows(only1, cols1, "акте 1", -1.0))
    result.extend(_fns_missing_rows(only2, cols2, "акте 2", 1.0))
    return result


def create_fns_discrepancy_protocol(output, rows1, rows2, cols1, cols2, differences, only1, only2, matches):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    t1, t2 = totals(rows1, cols1), totals(rows2, cols2)
    total_mode, debit_diff, credit_diff = reconcile_total_diffs(t1, t2)
    protocol_rows = _fns_discrepancy_rows(differences, only1, only2, cols1, cols2)
    has_disagreement = bool(protocol_rows)
    now = datetime.now()

    wb = Workbook()
    ws = wb.active
    ws.title = "ФНС_ПРОТОКОЛ"

    metadata = [
        ["ПРОТОКОЛ РАСХОЖДЕНИЙ К АКТУ СВЕРКИ ВЗАИМНЫХ РАСЧЕТОВ", ""],
        ["Основа структуры", FNS_RECONCILIATION_ORDER],
        ["Версия формата (ВерсФорм)", FNS_FORMAT_VERSION],
        ["Часть формата", FNS_FORMAT_PART],
        ["КНД ответного титула", FNS_RECIPIENT_KND],
        ["Признак наличия разногласий (ПризнРазн)", "2 — есть разногласия" if has_disagreement else "1 — нет разногласий"],
        ["Версия программы (ВерсПрог)", f"{APP_NAME} {__version__}"],
        ["Дата формирования (ДатаИнфПол)", now.strftime("%d.%m.%Y")],
        ["Время формирования (ВрИнфПол)", now.strftime("%H:%M:%S")],
        ["Режим сопоставления Д/К", "Зеркальный" if total_mode == "mirror" else "Прямой"],
        ["Итого расхождение дебет (ИтогоРасхДеб)", debit_diff if debit_diff is not None else 0.0],
        ["Итого расхождение кредит (ИтогоРасхКр)", credit_diff if credit_diff is not None else 0.0],
        ["Совпавших документов", len(matches)],
        ["Примечание", "Excel-представление реквизитов ответного титула ФНС. Для юридически значимого ЭДО официальный формат — XML по приказу ЕД-7-26/405@."],
    ]
    for row in metadata:
        ws.append(row)

    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["B1"].fill = PatternFill("solid", fgColor="17365D")
    ws.merge_cells("A1:B1")
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 43
    ws.column_dimensions["B"].width = 75
    ws.freeze_panes = "A2"

    details = wb.create_sheet("ДокСв_СвОпер")
    headers = [
        "НомСтр", "НаимДок", "НомДок", "ДатаДок", "ИтогоРасх",
        "ДатаОпер", "НаимОпер", "СумДебет", "СумКредит", "ДопИнф",
    ]
    details.append(headers)
    if protocol_rows:
        for idx, item in enumerate(protocol_rows, 1):
            details.append([
                idx,
                str(item.get("НаимДок") or "Документ")[:100],
                str(item.get("НомДок") or "б/н")[:50],
                item.get("ДатаДок") or now.strftime("%d.%m.%Y"),
                float(item.get("ИтогоРасх") or 0.0),
                item.get("ДатаОпер") or item.get("ДатаДок") or now.strftime("%d.%m.%Y"),
                str(item.get("НаимОпер") or "Расхождение")[:500],
                item.get("СумДебет"),
                item.get("СумКредит"),
                str(item.get("ДопИнф") or "")[:500],
            ])
    else:
        details.append([1, "—", "б/н", now.strftime("%d.%m.%Y"), 0.0, now.strftime("%d.%m.%Y"), "Разногласия не выявлены", None, None, "ПризнРазн = 1"])

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in (5, 8, 9):
        for cell in details.iter_cols(min_col=col, max_col=col, min_row=2):
            for value_cell in cell:
                if isinstance(value_cell.value, (int, float)):
                    value_cell.number_format = "#,##0.00"
    widths = [10, 24, 20, 14, 16, 14, 48, 16, 16, 55]
    for idx, width in enumerate(widths, 1):
        details.column_dimensions[get_column_letter(idx)].width = width
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions

    notes = wb.create_sheet("СПРАВКА_ФНС")
    notes.append(["Код элемента", "Назначение в формате ФНС"])
    notes_rows = [
        ["ПризнРазн", "1 — разногласий нет; 2 — разногласия есть"],
        ["НаимДок / НомДок / ДатаДок", "Реквизиты документа (таблица 7.6 формата)"],
        ["ИтогоРасх", "Итог расхождения по документу (таблица 7.6)"],
        ["НомСтр / ДатаОпер / НаимОпер", "Реквизиты операции (таблица 7.7)"],
        ["СумДебет / СумКредит", "Суммы операции по данным принимающей стороны"],
        ["ИтогоРасхДеб / ИтогоРасхКр", "Итоговые расхождения по состоянию расчетов (таблица 7.4)"],
    ]
    for row in notes_rows:
        notes.append(row)
    for cell in notes[1]:
        cell.fill = header_fill
        cell.font = header_font
    notes.column_dimensions["A"].width = 34
    notes.column_dimensions["B"].width = 90
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(output)
'''

once(
    "def get_android_activity():",
    FNS_FUNCTIONS + "\n\ndef get_android_activity():",
    "генератор протокола ФНС",
)

once(
    '''        self.save_button = ModernButton("СОХРАНИТЬ EXCEL-ОТЧЕТ", tone="neutral")
        self.save_button.bind(on_release=self.save_report)
        self.save_button.disabled = True
        actions.add_widget(self.save_button)''',
    '''        self.save_button = ModernButton("СОХРАНИТЬ EXCEL-ОТЧЕТ", tone="neutral")
        self.save_button.bind(on_release=self.save_report)
        self.save_button.disabled = True
        actions.add_widget(self.save_button)

        self.fns_button = ModernButton("ПРОТОКОЛ РАСХОЖДЕНИЙ • ФНС", tone="primary")
        self.fns_button.bind(on_release=self.save_fns_protocol)
        self.fns_button.disabled = True
        actions.add_widget(self.fns_button)''',
    "кнопка протокола ФНС",
)

once(
    '''            self.last_result = result
            self.save_button.disabled = False''',
    '''            self.last_result = result
            self.save_button.disabled = False
            self.fns_button.disabled = False''',
    "активация кнопки ФНС",
)

FNS_METHODS = r'''    def save_fns_protocol(self, *_):
        if not self.last_result:
            self.error("Сначала выполните сравнение актов.")
            return
        create_document(self.on_save_fns_uri, "Протокол_расхождений_ФНС.xlsx")

    def on_save_fns_uri(self, uri):
        try:
            differences, only1, only2, matches = self.last_result
            fd, path = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            try:
                create_fns_discrepancy_protocol(
                    path,
                    self.rows1, self.rows2,
                    self.cols1, self.cols2,
                    differences, only1, only2, matches,
                )
                copy_to_uri(path, uri)
            finally:
                if os.path.exists(path):
                    os.unlink(path)
            self.popup("Готово", "Протокол расхождений по структуре ФНС сохранен в Excel.", success=True)
        except Exception as exc:
            self.error(str(exc))

'''

once(
    "    def save_report(self, *_):",
    FNS_METHODS + "    def save_report(self, *_):",
    "сохранение протокола ФНС",
)

once(
    'actions.add_widget(text_label("XLS / XLSX / CSV • автокодировка • умное сопоставление • контроль дублей.", size=11, color=MUTED, height=38, valign="top"))',
    'actions.add_widget(text_label("XLS / XLSX / CSV • автосверка • протокол расхождений по структуре ФНС в Excel.", size=11, color=MUTED, height=38, valign="top"))',
    "подпись протокола",
)

path.write_text(text, encoding="utf-8")
print("FNS protocol patch applied: discrepancy protocol Excel + v1.5.0")
