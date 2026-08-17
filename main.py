import os
import posixpath
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kivy.app import App
from kivy.clock import Clock
from kivy.core.window import Window
from kivy.graphics import Color, Line, RoundedRectangle
from kivy.metrics import dp, sp
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.gridlayout import GridLayout
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.scrollview import ScrollView
from kivy.uix.textinput import TextInput
from jnius import autoclass

__version__ = "1.1.0"
APP_NAME = "Сверка PRO"

# Modern light theme
BG = (0.955, 0.965, 0.985, 1)
SURFACE = (1, 1, 1, 1)
SURFACE_SOFT = (0.965, 0.975, 0.995, 1)
TEXT = (0.075, 0.095, 0.14, 1)
MUTED = (0.39, 0.43, 0.51, 1)
BORDER = (0.865, 0.89, 0.94, 1)
PRIMARY = (0.16, 0.38, 0.90, 1)
PRIMARY_PRESSED = (0.11, 0.29, 0.72, 1)
SUCCESS = (0.07, 0.62, 0.34, 1)
SUCCESS_PRESSED = (0.05, 0.48, 0.27, 1)
DANGER = (0.91, 0.28, 0.25, 1)
WHITE = (1, 1, 1, 1)

XML_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def ntext(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower().replace("ё", "е"))


def nheader(value):
    return re.sub(r"[^a-zа-я0-9]+", "", ntext(value))


def number(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("\xa0", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"(руб\.?|₽|р\.)$", "", text, flags=re.I)
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(text)
    except Exception:
        return None


def date_key(value):
    if value is None or str(value).strip() == "":
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", text)
    if m:
        d, mo, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"
    return ntext(value)


def equal(a, b, tolerance=0.01):
    if a is None and b is None:
        return True
    if ntext(a) == ntext(b):
        return True
    na, nb = number(a), number(b)
    return na is not None and nb is not None and abs(na - nb) <= tolerance


ALIASES = {
    "date": ["дата", "дата операции", "дата документа", "датадокумента", "date"],
    "document": [
        "документ", "номер документа", "номердокумента", "№ документа", "номер",
        "док", "основание", "счет фактура", "счет-фактура", "накладная", "акт",
    ],
    "debit": ["дебет", "дебет оборот", "оборот дебет", "дебетовый оборот", "приход"],
    "credit": ["кредит", "кредит оборот", "оборот кредит", "кредитовый оборот", "расход"],
    "balance": ["сальдо", "остаток", "баланс", "конечное сальдо", "конечный остаток"],
    "counterparty": ["контрагент", "поставщик", "покупатель", "организация", "наименование", "партнер"],
}


def detect_column(headers, kind):
    best = None
    score_best = 0
    aliases = [nheader(x) for x in ALIASES[kind]]
    for h in headers:
        nh = nheader(h)
        score = 0
        for alias in aliases:
            if nh == alias:
                score = max(score, 100)
            elif alias and alias in nh:
                score = max(score, 70)
            elif nh and nh in alias:
                score = max(score, 50)
        if score > score_best:
            score_best = score
            best = h
    return best


def make_headers(raw_headers):
    headers, used = [], {}
    for i, h in enumerate(raw_headers or []):
        name = str(h).strip() if h is not None else ""
        name = name or f"Колонка {i + 1}"
        used[name] = used.get(name, 0) + 1
        if used[name] > 1:
            name = f"{name}_{used[name]}"
        headers.append(name)
    return headers


def choose_header_row(matrix):
    """Pick the most table-like row, useful for 1C exports with title rows above the table."""
    if not matrix:
        return 0
    best_idx = 0
    best_score = -1
    for idx, raw in enumerate(matrix[:40]):
        non_empty = sum(1 for x in raw if x is not None and str(x).strip())
        if non_empty == 0:
            continue
        headers = make_headers(raw)
        recognized = sum(1 for kind in ALIASES if detect_column(headers, kind))
        score = recognized * 100 + min(non_empty, 30)
        if score > best_score:
            best_score = score
            best_idx = idx
    return best_idx


def matrix_to_table(matrix, sheet_name):
    if not matrix:
        raise ValueError("В XLSX не найден непустой лист с таблицей.")
    header_idx = choose_header_row(matrix)
    headers = make_headers(matrix[header_idx])
    data = []
    for raw in matrix[header_idx + 1:]:
        if any(x is not None and str(x).strip() for x in raw):
            data.append({headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))})
    if not data:
        raise ValueError("В XLSX не найдено строк данных после заголовка таблицы.")
    return headers, data, sheet_name


def read_csv(path):
    import csv

    last = None
    for enc in ("utf-8-sig", "cp1251", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
                except Exception:
                    dialect = csv.excel
                matrix = list(csv.reader(f, dialect))
            if not matrix:
                raise ValueError("CSV-файл пуст.")
            return matrix_to_table(matrix, "CSV")
        except Exception as exc:
            last = exc
    raise last or ValueError("Не удалось прочитать CSV.")


def read_xlsx_openpyxl(path):
    errors = []
    for read_only in (True, False):
        try:
            wb = load_workbook(path, read_only=read_only, data_only=True)
            try:
                for ws in wb.worksheets:
                    matrix = [list(row) for row in ws.iter_rows(values_only=True)]
                    if any(any(x is not None and str(x).strip() for x in row) for row in matrix):
                        return matrix_to_table(matrix, ws.title)
            finally:
                wb.close()
        except Exception as exc:
            errors.append(exc)
    if errors:
        raise errors[-1]
    raise ValueError("Не удалось прочитать XLSX.")


def _find_case_insensitive(names, wanted):
    wanted_lower = wanted.lower()
    for name in names:
        if name.lower() == wanted_lower:
            return name
    return None


def _parse_shared_strings(zf, names):
    path = _find_case_insensitive(names, "xl/sharedStrings.xml")
    if not path:
        return []
    root = ET.fromstring(zf.read(path))
    strings = []
    for si in root.findall(".//main:si", XML_NS):
        strings.append("".join((node.text or "") for node in si.findall(".//main:t", XML_NS)))
    return strings


def _parse_date_style_indexes(zf, names):
    path = _find_case_insensitive(names, "xl/styles.xml")
    if not path:
        return set()
    try:
        root = ET.fromstring(zf.read(path))
        custom_formats = {}
        numfmts = root.find("main:numFmts", XML_NS)
        if numfmts is not None:
            for fmt in numfmts.findall("main:numFmt", XML_NS):
                try:
                    custom_formats[int(fmt.attrib.get("numFmtId", "0"))] = fmt.attrib.get("formatCode", "")
                except Exception:
                    pass

        builtin_date_ids = set(range(14, 23)) | {27, 30, 36, 45, 46, 47, 50, 57}
        date_styles = set()
        cell_xfs = root.find("main:cellXfs", XML_NS)
        if cell_xfs is None:
            return date_styles
        for idx, xf in enumerate(cell_xfs.findall("main:xf", XML_NS)):
            try:
                fmt_id = int(xf.attrib.get("numFmtId", "0"))
            except Exception:
                fmt_id = 0
            is_date = fmt_id in builtin_date_ids
            if not is_date and fmt_id in custom_formats:
                code = re.sub(r'"[^"]*"|\\.', "", custom_formats[fmt_id].lower())
                is_date = any(token in code for token in ("yy", "dd", "mm", "hh", "ss"))
            if is_date:
                date_styles.add(idx)
        return date_styles
    except Exception:
        return set()


def _sheet_paths(zf, names):
    workbook_path = _find_case_insensitive(names, "xl/workbook.xml")
    if not workbook_path:
        return []
    wb_root = ET.fromstring(zf.read(workbook_path))

    rels_path = _find_case_insensitive(names, "xl/_rels/workbook.xml.rels")
    rel_map = {}
    if rels_path:
        rel_root = ET.fromstring(zf.read(rels_path))
        for rel in rel_root.findall(".//pkgrel:Relationship", XML_NS):
            rel_map[rel.attrib.get("Id")] = rel.attrib.get("Target", "")

    result = []
    for sheet in wb_root.findall(".//main:sheets/main:sheet", XML_NS):
        sheet_name = sheet.attrib.get("name", "Лист")
        rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rel_map.get(rel_id, "")
        if not target:
            continue
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = posixpath.normpath(posixpath.join("xl", target))
        actual = _find_case_insensitive(names, path)
        if actual:
            result.append((sheet_name, actual))

    if not result:
        for path in sorted(names):
            if path.lower().startswith("xl/worksheets/") and path.lower().endswith(".xml"):
                result.append((Path(path).stem, path))
    return result


def _column_index(cell_ref):
    m = re.match(r"([A-Z]+)", str(cell_ref).upper())
    if not m:
        return 0
    value = 0
    for ch in m.group(1):
        value = value * 26 + ord(ch) - 64
    return value - 1


def _excel_date(value):
    try:
        serial = float(value)
        dt = datetime(1899, 12, 30) + timedelta(days=serial)
        if dt.time().hour == 0 and dt.time().minute == 0 and dt.time().second == 0:
            return dt.date()
        return dt
    except Exception:
        return value


def _cell_value(cell, shared_strings, date_styles):
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join((node.text or "") for node in cell.findall(".//main:is//main:t", XML_NS))

    value_node = cell.find("main:v", XML_NS)
    raw = value_node.text if value_node is not None else ""

    if cell_type == "s":
        try:
            idx = int(raw)
            return shared_strings[idx] if 0 <= idx < len(shared_strings) else ""
        except Exception:
            return ""
    if cell_type == "b":
        return raw == "1"
    if cell_type == "str":
        return raw

    try:
        style_index = int(cell.attrib.get("s", "-1"))
    except Exception:
        style_index = -1
    if raw and style_index in date_styles:
        return _excel_date(raw)

    if raw != "":
        try:
            num = float(raw)
            return int(num) if num.is_integer() else num
        except Exception:
            pass
    return raw


def read_xlsx_fallback(path):
    """Read imperfect 1C/third-party XLSX packages without relying on sharedStrings.xml."""
    with zipfile.ZipFile(path, "r") as zf:
        names = zf.namelist()
        shared_strings = _parse_shared_strings(zf, names)
        date_styles = _parse_date_style_indexes(zf, names)
        sheets = _sheet_paths(zf, names)
        if not sheets:
            raise ValueError("В XLSX не удалось найти листы.")

        for sheet_name, sheet_path in sheets:
            root = ET.fromstring(zf.read(sheet_path))
            matrix = []
            for row in root.findall(".//main:sheetData/main:row", XML_NS):
                values = {}
                max_index = -1
                for cell in row.findall("main:c", XML_NS):
                    idx = _column_index(cell.attrib.get("r", ""))
                    values[idx] = _cell_value(cell, shared_strings, date_styles)
                    max_index = max(max_index, idx)
                if max_index >= 0:
                    matrix.append([values.get(i, None) for i in range(max_index + 1)])
            if matrix and any(any(x is not None and str(x).strip() for x in row) for row in matrix):
                return matrix_to_table(matrix, sheet_name)
    raise ValueError("В XLSX не найден непустой лист с таблицей.")


def read_xlsx(path):
    try:
        return read_xlsx_openpyxl(path)
    except Exception as primary_error:
        try:
            return read_xlsx_fallback(path)
        except Exception as fallback_error:
            raise ValueError(
                "Не удалось прочитать XLSX. Файл может быть поврежден или иметь нестандартную структуру. "
                f"Основной режим: {primary_error}. Резервный режим: {fallback_error}"
            )


def read_table(path):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        return read_csv(path)
    if ext == ".xlsx":
        return read_xlsx(path)
    raise ValueError("Поддерживаются XLSX и CSV. Файл XLS сначала сохраните как XLSX.")


def make_key(row, cols, occurrence):
    doc = ntext(row.get(cols.get("document"))) if cols.get("document") else ""
    dt = date_key(row.get(cols.get("date"))) if cols.get("date") else ""
    if doc or dt:
        return f"{doc}|{dt}"
    debit = number(row.get(cols.get("debit"))) if cols.get("debit") else None
    credit = number(row.get(cols.get("credit"))) if cols.get("credit") else None
    return f"amount|{debit}|{credit}|{occurrence}"


def build_index(rows, cols):
    result, count = {}, {}
    for i, row in enumerate(rows):
        base = make_key(row, cols, i)
        n = count.get(base, 0)
        count[base] = n + 1
        result[f"{base}||{n}"] = i
    return result


def compare(rows1, rows2, cols1, cols2, tolerance):
    idx1, idx2 = build_index(rows1, cols1), build_index(rows2, cols2)
    pairs = []
    for kind in ("date", "document", "debit", "credit", "balance", "counterparty"):
        c1, c2 = cols1.get(kind), cols2.get(kind)
        if c1 and c2:
            pairs.append((kind, c1, c2))

    differences, only1, only2, matches = [], [], [], []
    for key in dict.fromkeys(list(idx1) + list(idx2)):
        if key not in idx2:
            only1.append(rows1[idx1[key]].copy())
            continue
        if key not in idx1:
            only2.append(rows2[idx2[key]].copy())
            continue
        r1, r2 = rows1[idx1[key]], rows2[idx2[key]]
        changed = []
        for kind, c1, c2 in pairs:
            a, b = r1.get(c1), r2.get(c2)
            if not equal(a, b, tolerance):
                na, nb = number(a), number(b)
                delta = nb - na if na is not None and nb is not None else ""
                changed.append((kind, c1, c2, a, b, delta))
        if changed:
            for kind, c1, c2, a, b, delta in changed:
                differences.append([key, kind.upper(), c1, c2, a, b, delta, "РАЗЛИЧИЕ"])
        else:
            matches.append([key, "Совпадает"])
    return differences, only1, only2, matches


def totals(rows, cols):
    out = {}
    for kind in ("debit", "credit", "balance"):
        col = cols.get(kind)
        if not col:
            out[kind] = None
            continue
        values = [number(r.get(col)) for r in rows]
        values = [x for x in values if x is not None]
        out[kind] = sum(values) if values else 0.0
    return out


def money(value):
    if value is None:
        return "—"
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def autosize(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def create_report(output, rows1, rows2, cols1, cols2, differences, only1, only2, matches, name1, name2):
    t1, t2 = totals(rows1, cols1), totals(rows2, cols2)

    def diff(kind):
        return t2[kind] - t1[kind] if t1[kind] is not None and t2[kind] is not None else None

    wb = Workbook()
    ws = wb.active
    ws.title = "ИТОГ"
    summary = [
        ["СВЕРКА PRO — ИТОГ", ""], ["Акт 1", name1], ["Акт 2", name2],
        ["Строк в акте 1", len(rows1)], ["Строк в акте 2", len(rows2)],
        ["Строк с расхождениями", len(set(x[0] for x in differences))],
        ["Всего отдельных расхождений", len(differences)], ["Только в акте 1", len(only1)],
        ["Только в акте 2", len(only2)], ["Совпадения", len(matches)],
        ["ИТОГО ДЕБЕТ — акт 1", t1["debit"]], ["ИТОГО ДЕБЕТ — акт 2", t2["debit"]],
        ["РАЗНИЦА ДЕБЕТ", diff("debit")], ["ИТОГО КРЕДИТ — акт 1", t1["credit"]],
        ["ИТОГО КРЕДИТ — акт 2", t2["credit"]], ["РАЗНИЦА КРЕДИТ", diff("credit")],
        ["САЛЬДО — акт 1", t1["balance"]], ["САЛЬДО — акт 2", t2["balance"]],
        ["РАЗНИЦА САЛЬДО", diff("balance")],
    ]
    for row in summary:
        ws.append(row)

    wsd = wb.create_sheet("РАЗЛИЧИЯ")
    wsd.append(["Ключ", "Поле", "Колонка акта 1", "Колонка акта 2", "Акт 1", "Акт 2", "Разница", "Статус"])
    for row in differences:
        wsd.append(row)

    def write_rows(title, data):
        sheet = wb.create_sheet(title)
        if not data:
            sheet.append(["Нет данных"])
            return
        headers = list(data[0].keys())
        sheet.append(headers)
        for row in data:
            sheet.append([row.get(h, "") for h in headers])

    write_rows("ТОЛЬКО АКТ 1", only1)
    write_rows("ТОЛЬКО АКТ 2", only2)
    wsm = wb.create_sheet("СОВПАДЕНИЯ")
    wsm.append(["Ключ", "Статус"])
    for row in matches:
        wsm.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        autosize(sheet)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"

    for row in range(2, wsd.max_row + 1):
        wsd.cell(row, 5).fill = PatternFill("solid", fgColor="FECACA")
        wsd.cell(row, 6).fill = PatternFill("solid", fgColor="DCFCE7")
        wsd.cell(row, 7).fill = PatternFill("solid", fgColor="FEF3C7")
    wb.save(output)


def get_android_activity():
    return autoclass("org.kivy.android.PythonActivity").mActivity


def open_document(callback):
    Intent = autoclass("android.content.Intent")
    intent = Intent(Intent.ACTION_OPEN_DOCUMENT)
    intent.addCategory(Intent.CATEGORY_OPENABLE)
    intent.setType("*/*")
    intent.putExtra(Intent.EXTRA_MIME_TYPES, [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/csv",
        "text/comma-separated-values",
    ])
    get_android_activity().startActivityForResult(intent, 7001)
    App.get_running_app().pending_open_callback = callback


def create_document(callback):
    Intent = autoclass("android.content.Intent")
    intent = Intent(Intent.ACTION_CREATE_DOCUMENT)
    intent.addCategory(Intent.CATEGORY_OPENABLE)
    intent.setType("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    intent.putExtra(Intent.EXTRA_TITLE, "Отчет_сравнения_актов.xlsx")
    get_android_activity().startActivityForResult(intent, 7002)
    App.get_running_app().pending_save_callback = callback


def uri_suffix(uri):
    resolver = get_android_activity().getContentResolver()
    mime = str(resolver.getType(uri) or "").lower()
    if "spreadsheetml" in mime or "xlsx" in mime:
        return ".xlsx"
    if "csv" in mime:
        return ".csv"

    cursor = None
    try:
        OpenableColumns = autoclass("android.provider.OpenableColumns")
        cursor = resolver.query(uri, None, None, None, None)
        if cursor and cursor.moveToFirst():
            idx = cursor.getColumnIndex(OpenableColumns.DISPLAY_NAME)
            if idx >= 0:
                ext = Path(str(cursor.getString(idx))).suffix.lower()
                if ext in (".xlsx", ".csv"):
                    return ext
    except Exception:
        pass
    finally:
        if cursor:
            try:
                cursor.close()
            except Exception:
                pass
    return ".xlsx"


def uri_to_tempfile(uri):
    resolver = get_android_activity().getContentResolver()
    stream = resolver.openInputStream(uri)
    if stream is None:
        raise ValueError("Android не смог открыть выбранный файл.")
    fd, path = tempfile.mkstemp(suffix=uri_suffix(uri))
    os.close(fd)
    try:
        with open(path, "wb") as out:
            buf = bytearray(64 * 1024)
            while True:
                count = stream.read(buf)
                if count <= 0:
                    break
                out.write(bytes(buf[:count]))
    finally:
        stream.close()
    return path


def copy_to_uri(local_path, uri):
    resolver = get_android_activity().getContentResolver()
    stream = resolver.openOutputStream(uri)
    if stream is None:
        raise ValueError("Android не смог создать выбранный файл.")
    try:
        with open(local_path, "rb") as src:
            while True:
                chunk = src.read(64 * 1024)
                if not chunk:
                    break
                stream.write(chunk)
        stream.flush()
    finally:
        stream.close()


class Card(BoxLayout):
    def __init__(self, radius=18, **kwargs):
        super().__init__(orientation="vertical", spacing=dp(10), padding=dp(16), size_hint_y=None, **kwargs)
        self.radius = dp(radius)
        self.bind(minimum_height=self.setter("height"))
        with self.canvas.before:
            self._bg_color = Color(*SURFACE)
            self._bg_rect = RoundedRectangle(pos=self.pos, size=self.size, radius=[self.radius])
            self._border_color = Color(*BORDER)
            self._border_line = Line(
                rounded_rectangle=(self.x, self.y, self.width, self.height, self.radius),
                width=1,
            )
        self.bind(pos=self._sync_canvas, size=self._sync_canvas)

    def _sync_canvas(self, *_):
        self._bg_rect.pos = self.pos
        self._bg_rect.size = self.size
        self._border_line.rounded_rectangle = (self.x, self.y, self.width, self.height, self.radius)


class ModernButton(Button):
    def __init__(self, text, tone="primary", **kwargs):
        super().__init__(
            text=text,
            size_hint_y=None,
            height=dp(52),
            background_normal="",
            background_down="",
            background_disabled_normal="",
            background_color=(0, 0, 0, 0),
            bold=True,
            font_size=sp(15),
            **kwargs,
        )
        self.tone = tone
        self._radius = dp(15)
        with self.canvas.before:
            self._fill = Color(*PRIMARY)
            self._rect = RoundedRectangle(pos=self.pos, size=self.size, radius=[self._radius])
        self.bind(pos=self._sync_canvas, size=self._sync_canvas, state=self._refresh, disabled=self._refresh)
        self._refresh()

    def _palette(self):
        if self.disabled:
            return (0.82, 0.84, 0.88, 1), (0.48, 0.51, 0.57, 1)
        if self.tone == "success":
            bg = SUCCESS_PRESSED if self.state == "down" else SUCCESS
            return bg, WHITE
        if self.tone == "neutral":
            bg = (0.90, 0.92, 0.96, 1) if self.state == "down" else SURFACE_SOFT
            return bg, TEXT
        bg = PRIMARY_PRESSED if self.state == "down" else PRIMARY
        return bg, WHITE

    def _refresh(self, *_):
        bg, fg = self._palette()
        self._fill.rgba = bg
        self.color = fg

    def _sync_canvas(self, *_):
        self._rect.pos = self.pos
        self._rect.size = self.size


def text_label(text, size=13, color=TEXT, bold=False, height=28, markup=False, valign="middle"):
    label = Label(
        text=text,
        color=color,
        font_size=sp(size),
        bold=bold,
        size_hint_y=None,
        height=dp(height),
        halign="left",
        valign=valign,
        markup=markup,
    )
    label.bind(size=lambda inst, value: setattr(inst, "text_size", (value[0], None)))
    return label


class MainScreen(BoxLayout):
    def __init__(self, **kwargs):
        super().__init__(orientation="vertical", **kwargs)
        self.rows1 = self.rows2 = None
        self.headers1, self.headers2 = [], []
        self.cols1, self.cols2 = {}, {}
        self.last_result = None
        self.build_ui()

    def build_ui(self):
        scroll = ScrollView(do_scroll_x=False, bar_width=dp(4), scroll_type=["bars", "content"])
        root = BoxLayout(orientation="vertical", spacing=dp(14), padding=[dp(14), dp(14), dp(14), dp(24)], size_hint_y=None)
        root.bind(minimum_height=root.setter("height"))

        header = Card(radius=22)
        header.add_widget(text_label(APP_NAME, size=24, bold=True, height=38))
        header.add_widget(text_label("Сверяйте два акта и сохраняйте понятный Excel-отчет", size=13, color=MUTED, height=38, valign="top"))
        header.add_widget(text_label(f"Версия {__version__}  •  Android", size=11, color=PRIMARY, bold=True, height=22))
        root.add_widget(header)

        files = Card()
        files.add_widget(text_label("Файлы", size=17, bold=True, height=30))
        files.add_widget(text_label("Выберите два акта в XLSX или CSV", size=12, color=MUTED, height=24))

        self.file1_button = ModernButton("Выбрать АКТ 1")
        self.file1_button.bind(on_release=self.choose1)
        files.add_widget(self.file1_button)
        self.file1_label = text_label("Файл 1 не выбран", size=12, color=MUTED, height=26)
        files.add_widget(self.file1_label)

        self.file2_button = ModernButton("Выбрать АКТ 2")
        self.file2_button.bind(on_release=self.choose2)
        files.add_widget(self.file2_button)
        self.file2_label = text_label("Файл 2 не выбран", size=12, color=MUTED, height=26)
        files.add_widget(self.file2_label)
        root.add_widget(files)

        fields = Card()
        fields.add_widget(text_label("Распознанные поля", size=17, bold=True, height=30))
        self.fields_grid = GridLayout(cols=2, spacing=dp(8), size_hint_y=None)
        self.fields_grid.bind(minimum_height=self.fields_grid.setter("height"))
        fields.add_widget(self.fields_grid)
        root.add_widget(fields)
        self.update_fields()

        settings = Card()
        settings.add_widget(text_label("Настройки", size=17, bold=True, height=30))
        tol_row = BoxLayout(size_hint_y=None, height=dp(48), spacing=dp(10))
        tol_row.add_widget(text_label("Допуск, руб.", size=14, height=48))
        self.tolerance = TextInput(
            text="0.01",
            multiline=False,
            input_filter="float",
            size_hint_x=0.42,
            font_size=sp(15),
            foreground_color=TEXT,
            cursor_color=PRIMARY,
            background_normal="",
            background_active="",
            background_color=SURFACE_SOFT,
            padding=[dp(12), dp(12), dp(12), dp(12)],
        )
        tol_row.add_widget(self.tolerance)
        settings.add_widget(tol_row)
        root.add_widget(settings)

        result = Card()
        result.add_widget(text_label("Результат", size=17, bold=True, height=30))
        self.result_label = Label(
            text="Выберите оба акта. После сравнения здесь появятся расхождения и итоговые суммы.",
            color=MUTED,
            font_size=sp(13),
            size_hint_y=None,
            height=dp(132),
            halign="left",
            valign="top",
            markup=True,
        )
        self.result_label.bind(size=lambda inst, value: setattr(inst, "text_size", (value[0], None)))
        result.add_widget(self.result_label)
        root.add_widget(result)

        actions = Card()
        self.compare_button = ModernButton("СРАВНИТЬ АКТЫ", tone="success")
        self.compare_button.bind(on_release=self.run_compare)
        actions.add_widget(self.compare_button)

        self.save_button = ModernButton("СОХРАНИТЬ EXCEL-ОТЧЕТ", tone="neutral")
        self.save_button.bind(on_release=self.save_report)
        self.save_button.disabled = True
        actions.add_widget(self.save_button)
        actions.add_widget(text_label("Поддерживаются XLSX и CSV. Нестандартные XLSX из 1С читаются в резервном режиме.", size=11, color=MUTED, height=38, valign="top"))
        root.add_widget(actions)

        scroll.add_widget(root)
        self.add_widget(scroll)

    def choose1(self, *_):
        open_document(self.open1)

    def choose2(self, *_):
        open_document(self.open2)

    def load_file(self, uri, target):
        path = None
        try:
            path = uri_to_tempfile(uri)
            headers, rows, sheet = read_table(path)
            cols = {key: detect_column(headers, key) for key in ALIASES}
            if target == 1:
                self.rows1, self.headers1, self.cols1 = rows, headers, cols
                self.file1_label.text = f"АКТ 1 загружен: {len(rows)} строк • {sheet}"
                self.file1_label.color = SUCCESS
            else:
                self.rows2, self.headers2, self.cols2 = rows, headers, cols
                self.file2_label.text = f"АКТ 2 загружен: {len(rows)} строк • {sheet}"
                self.file2_label.color = SUCCESS
            self.update_fields()
        except Exception as exc:
            self.error(str(exc))
        finally:
            if path and os.path.exists(path):
                try:
                    os.unlink(path)
                except Exception:
                    pass

    def open1(self, uri):
        self.load_file(uri, 1)

    def open2(self, uri):
        self.load_file(uri, 2)

    def field_label(self, title, value, primary=False):
        color = PRIMARY if primary and value != "—" else (TEXT if value != "—" else MUTED)
        label = Label(
            text=f"[b]{title}[/b]\n{value}",
            markup=True,
            color=color,
            font_size=sp(11.5),
            size_hint_y=None,
            height=dp(58),
            halign="left",
            valign="middle",
        )
        label.bind(size=lambda inst, val: setattr(inst, "text_size", (val[0], None)))
        return label

    def update_fields(self):
        self.fields_grid.clear_widgets()
        if self.rows1 is None and self.rows2 is None:
            self.fields_grid.add_widget(text_label("Поля появятся после загрузки файлов", size=12, color=MUTED, height=36))
            self.fields_grid.add_widget(Label(text="", size_hint_y=None, height=dp(36)))
            return
        for kind, title in (
            ("date", "Дата"),
            ("document", "Документ"),
            ("debit", "Дебет"),
            ("credit", "Кредит"),
            ("balance", "Сальдо"),
            ("counterparty", "Контрагент"),
        ):
            c1 = self.cols1.get(kind) or "—"
            c2 = self.cols2.get(kind) or "—"
            self.fields_grid.add_widget(self.field_label(f"{title} • Акт 1", c1, primary=True))
            self.fields_grid.add_widget(self.field_label(f"{title} • Акт 2", c2))

    def run_compare(self, *_):
        if self.rows1 is None or self.rows2 is None:
            self.error("Сначала выберите оба акта сверки.")
            return
        if not self.cols1.get("document") and not self.cols1.get("date"):
            self.error("В акте 1 не найдено поле «Документ» или «Дата».")
            return
        if not self.cols2.get("document") and not self.cols2.get("date"):
            self.error("В акте 2 не найдено поле «Документ» или «Дата».")
            return

        try:
            tolerance = float(self.tolerance.text.replace(",", "."))
            result = compare(self.rows1, self.rows2, self.cols1, self.cols2, tolerance)
            self.last_result = result
            self.save_button.disabled = False
            differences, only1, only2, matches = result
            t1, t2 = totals(self.rows1, self.cols1), totals(self.rows2, self.cols2)
            debit = t2["debit"] - t1["debit"] if t1["debit"] is not None and t2["debit"] is not None else None
            credit = t2["credit"] - t1["credit"] if t1["credit"] is not None and t2["credit"] is not None else None
            diff_rows = len(set(x[0] for x in differences))
            self.result_label.color = TEXT
            self.result_label.text = (
                f"[b]Расхождения:[/b] {diff_rows}\n"
                f"Отдельных отличий: {len(differences)}\n"
                f"Только в акте 1: {len(only1)}\n"
                f"Только в акте 2: {len(only2)}\n"
                f"Совпадений: {len(matches)}\n\n"
                f"[b]Разница ДЕБЕТ:[/b] {money(debit)} руб.\n"
                f"[b]Разница КРЕДИТ:[/b] {money(credit)} руб."
            )
        except Exception as exc:
            self.error(str(exc))

    def save_report(self, *_):
        if not self.last_result:
            self.error("Сначала выполните сравнение.")
            return
        create_document(self.on_save_uri)

    def on_save_uri(self, uri):
        try:
            differences, only1, only2, matches = self.last_result
            fd, path = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            try:
                create_report(
                    path, self.rows1, self.rows2, self.cols1, self.cols2,
                    differences, only1, only2, matches, "Акт 1", "Акт 2"
                )
                copy_to_uri(path, uri)
            finally:
                if os.path.exists(path):
                    os.unlink(path)
            self.popup("Готово", "Excel-отчет успешно сохранен.", success=True)
        except Exception as exc:
            self.error(str(exc))

    def error(self, text):
        self.popup("Ошибка", text, success=False)

    def popup(self, title, message, success=False):
        box = BoxLayout(orientation="vertical", padding=dp(16), spacing=dp(14))
        label = Label(
            text=message,
            color=WHITE,
            font_size=sp(14),
            halign="left",
            valign="top",
        )
        label.bind(size=lambda inst, val: setattr(inst, "text_size", (val[0], None)))
        box.add_widget(label)

        button = ModernButton("OK", tone="success" if success else "primary")
        box.add_widget(button)

        pop = Popup(
            title=title,
            title_size=sp(18),
            content=box,
            size_hint=(0.90, 0.43),
            auto_dismiss=False,
            separator_color=SUCCESS if success else DANGER,
        )
        button.bind(on_release=pop.dismiss)
        pop.open()


class AndroidApp(App):
    def build(self):
        Window.clearcolor = BG
        self.pending_open_callback = None
        self.pending_save_callback = None
        from android import activity
        self.android_activity = activity
        activity.bind(on_activity_result=self.on_activity_result)
        return MainScreen()

    def on_stop(self):
        try:
            self.android_activity.unbind(on_activity_result=self.on_activity_result)
        except Exception:
            pass

    def on_activity_result(self, request_code, result_code, intent):
        if result_code != -1 or intent is None:
            return
        uri = intent.getData()
        if uri is None:
            return
        if request_code == 7001 and self.pending_open_callback:
            callback = self.pending_open_callback
            self.pending_open_callback = None
            Clock.schedule_once(lambda dt: callback(uri), 0)
        elif request_code == 7002 and self.pending_save_callback:
            callback = self.pending_save_callback
            self.pending_save_callback = None
            Clock.schedule_once(lambda dt: callback(uri), 0)


if __name__ == "__main__":
    AndroidApp().run()
